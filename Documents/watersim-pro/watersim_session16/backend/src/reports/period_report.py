#!/usr/bin/env python3
"""
WaterSim Pro — Period (historian) report
Reads JSON from stdin, writes PDF or XLSX bytes to stdout.

Payload (from periodReport.js):
  format        'pdf' | 'xlsx'
  title, org, generatedAt, generatedBy, from, to, bucket
  series[]      { tag, name, unit, area, kind, stats{min,max,avg,last,samples,availabilityPct},
                  points[[iso, avg, min, max], ...] }
  events[]      { triggeredAt, clearedAt, severity, state, rule, message, acknowledged }
  counts        { critical, warning, info }

Dependencies: reportlab, matplotlib (pdf); openpyxl (xlsx)
"""

import io
import json
import sys
from datetime import datetime


def parse_ts(s):
    return datetime.fromisoformat(s.replace('Z', '+00:00'))


def fmt_num(v, dp=2):
    if v is None:
        return '—'
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if abs(f) >= 1000:
        return f'{f:,.0f}'
    return f'{f:.{dp}f}'


def fmt_ts(s):
    if not s:
        return '—'
    return parse_ts(s).strftime('%d %b %Y %H:%M')


# ═══════════════════════════════════════════════════════════════════════════
# PDF
# ═══════════════════════════════════════════════════════════════════════════

def build_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                    HRFlowable, KeepTogether, Image as RLImage, PageBreak)
    from xml.sax.saxutils import escape
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates

    BLUE = colors.HexColor('#1E40AF'); CYAN = colors.HexColor('#0891B2'); GREY = colors.HexColor('#6B7280')
    RED = colors.HexColor('#DC2626'); AMBER = colors.HexColor('#D97706'); LGREY = colors.HexColor('#F3F4F6')
    MGREY = colors.HexColor('#E5E7EB')

    base = getSampleStyleSheet()
    S = {}
    def add(name, **kw):
        S[name] = ParagraphStyle(name, parent=base['Normal'], **kw)
    add('Title', fontSize=20, textColor=BLUE, fontName='Helvetica-Bold', spaceAfter=2)
    add('Sub', fontSize=10, textColor=GREY)
    add('H1', fontSize=13, textColor=BLUE, fontName='Helvetica-Bold', spaceBefore=12, spaceAfter=6)
    add('H2', fontSize=10.5, textColor=CYAN, fontName='Helvetica-Bold', spaceBefore=8, spaceAfter=3)
    add('Body', fontSize=9, leading=13)
    add('Small', fontSize=8, textColor=GREY)
    add('TH', fontSize=8, textColor=colors.white, fontName='Helvetica-Bold', alignment=TA_CENTER)
    add('TD', fontSize=8, alignment=TA_CENTER)
    add('TDL', fontSize=8, alignment=TA_LEFT)
    add('Cap', fontSize=8, textColor=GREY, alignment=TA_CENTER, fontName='Helvetica-Oblique', spaceAfter=6)

    W, H = A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=data.get('title') or 'Plant history report', author='WaterSim Pro')
    story = []
    esc = lambda s: escape(str(s if s is not None else ''))

    # ── Title block ──
    story.append(Paragraph(esc(data.get('title') or 'Plant history report'), S['Title']))
    story.append(Paragraph(esc(data.get('org') or ''), S['Sub']))
    story.append(Paragraph(
        f"{fmt_ts(data['from'])} → {fmt_ts(data['to'])} · resolution {esc(data.get('bucket'))} · "
        f"generated {fmt_ts(data['generatedAt'])}" + (f" by {esc(data['generatedBy'])}" if data.get('generatedBy') else ''),
        S['Sub']))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width='100%', thickness=1, color=MGREY))

    # ── Summary table ──
    story.append(Paragraph('Summary', S['H1']))
    head = ['Tag', 'Point', 'Unit', 'Min', 'Average', 'Max', 'Last', 'Availability']
    rows = [[Paragraph(h, S['TH']) for h in head]]
    for s in data['series']:
        st = s.get('stats') or {}
        rows.append([
            Paragraph(esc(s['tag']), S['TDL']), Paragraph(esc(s.get('name')), S['TDL']),
            Paragraph(esc(s.get('unit') or ''), S['TD']),
            Paragraph(fmt_num(st.get('min')), S['TD']), Paragraph(fmt_num(st.get('avg')), S['TD']),
            Paragraph(fmt_num(st.get('max')), S['TD']), Paragraph(fmt_num(st.get('last')), S['TD']),
            Paragraph(f"{fmt_num(st.get('availabilityPct'), 1)} %", S['TD']),
        ])
    t = Table(rows, colWidths=[28 * mm, 52 * mm, 14 * mm, 18 * mm, 18 * mm, 18 * mm, 18 * mm, 20 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), BLUE), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, LGREY]),
        ('GRID', (0, 0), (-1, -1), 0.4, MGREY), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(t)
    c = data.get('counts') or {}
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        f"Alarm events in the period: <b>{c.get('critical', 0)}</b> critical · "
        f"<b>{c.get('warning', 0)}</b> warning · <b>{c.get('info', 0)}</b> info", S['Body']))

    # ── One chart per tag ──
    story.append(Paragraph('Trends', S['H1']))
    for s in data['series']:
        pts = [p for p in s['points'] if p[1] is not None]
        if not pts:
            story.append(KeepTogether([Paragraph(f"{esc(s['tag'])} — {esc(s.get('name'))}", S['H2']),
                                       Paragraph('No good samples in this window.', S['Small'])]))
            continue
        xs = [parse_ts(p[0]) for p in pts]
        avg = [p[1] for p in pts]
        lo = [p[2] if p[2] is not None else p[1] for p in pts]
        hi = [p[3] if p[3] is not None else p[1] for p in pts]

        fig, ax = plt.subplots(figsize=(7.2, 2.3), dpi=150)
        ax.fill_between(xs, lo, hi, color='#93C5FD', alpha=0.35, linewidth=0, label='min–max')
        ax.plot(xs, avg, color='#1E40AF', linewidth=1.2, label='average')
        rmin, rmax = s.get('rangeMin'), s.get('rangeMax')
        if rmin is not None and rmax is not None and rmax > rmin:
            ax.axhline(rmin, color='#9CA3AF', linewidth=0.6, linestyle='--')
            ax.axhline(rmax, color='#9CA3AF', linewidth=0.6, linestyle='--')
        ax.set_ylabel(s.get('unit') or '', fontsize=8)
        ax.tick_params(labelsize=7)
        ax.grid(True, color='#E5E7EB', linewidth=0.5)
        for sp in ('top', 'right'):
            ax.spines[sp].set_visible(False)
        span_h = (xs[-1] - xs[0]).total_seconds() / 3600 if len(xs) > 1 else 1
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%H:%M' if span_h <= 48 else '%d %b'))
        ax.legend(fontsize=7, loc='upper right', frameon=False)
        fig.tight_layout()
        img = io.BytesIO()
        fig.savefig(img, format='png')
        plt.close(fig)
        img.seek(0)
        st = s.get('stats') or {}
        story.append(KeepTogether([
            Paragraph(f"{esc(s['tag'])} — {esc(s.get('name'))}", S['H2']),
            RLImage(img, width=W - 32 * mm, height=(W - 32 * mm) * 2.3 / 7.2),
            Paragraph(f"min {fmt_num(st.get('min'))} · avg {fmt_num(st.get('avg'))} · max {fmt_num(st.get('max'))} "
                      f"{esc(s.get('unit') or '')} · {st.get('samples', 0)} samples · "
                      f"{fmt_num(st.get('availabilityPct'), 1)} % available", S['Cap']),
        ]))

    # ── Alarm events ──
    events = data.get('events') or []
    if events:
        story.append(PageBreak())
        story.append(Paragraph('Alarm events in the period', S['H1']))
        head = ['Raised', 'Cleared', 'Severity', 'Rule', 'Message', 'Ack']
        rows = [[Paragraph(h, S['TH']) for h in head]]
        for e in events[:120]:
            rows.append([
                Paragraph(fmt_ts(e.get('triggeredAt')), S['TD']), Paragraph(fmt_ts(e.get('clearedAt')), S['TD']),
                Paragraph(esc(e.get('severity')), S['TD']), Paragraph(esc(e.get('rule')), S['TDL']),
                Paragraph(esc(e.get('message')), S['TDL']), Paragraph('yes' if e.get('acknowledged') else '', S['TD']),
            ])
        t = Table(rows, colWidths=[24 * mm, 24 * mm, 16 * mm, 44 * mm, 66 * mm, 10 * mm], repeatRows=1)
        style = [('BACKGROUND', (0, 0), (-1, 0), BLUE), ('GRID', (0, 0), (-1, -1), 0.4, MGREY),
                 ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, LGREY])]
        for i, e in enumerate(events[:120], start=1):
            col = RED if e.get('severity') == 'critical' else AMBER if e.get('severity') == 'warning' else GREY
            style.append(('TEXTCOLOR', (2, i), (2, i), col))
        t.setStyle(TableStyle(style))
        story.append(t)
        if len(events) > 120:
            story.append(Paragraph(f'… and {len(events) - 120} more (export the alarm CSV for the full list).', S['Small']))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(GREY)
        canvas.drawString(16 * mm, 10 * mm, f"WaterSim Pro · {data.get('org') or ''} · historian report")
        canvas.drawRightString(W - 16 * mm, 10 * mm, f'Page {doc_.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# XLSX
# ═══════════════════════════════════════════════════════════════════════════

def build_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import LineChart, Reference

    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='1E40AF')
    thin = Side(style='thin', color='E5E7EB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def header(ws, row, values):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = head_font; cell.fill = head_fill; cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # ── Summary ──
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = data.get('title') or 'Plant history report'; ws['A1'].font = Font(bold=True, size=14, color='1E40AF')
    ws['A2'] = data.get('org') or ''
    ws['A3'] = f"{fmt_ts(data['from'])} → {fmt_ts(data['to'])}  ·  resolution {data.get('bucket')}"
    ws['A4'] = f"Generated {fmt_ts(data['generatedAt'])}" + (f" by {data['generatedBy']}" if data.get('generatedBy') else '')
    header(ws, 6, ['Tag', 'Point', 'Area', 'Unit', 'Min', 'Average', 'Max', 'Last', 'Samples', 'Availability %'])
    for i, s in enumerate(data['series'], start=7):
        st = s.get('stats') or {}
        vals = [s['tag'], s.get('name'), s.get('area'), s.get('unit'), st.get('min'), st.get('avg'), st.get('max'),
                st.get('last'), st.get('samples'), st.get('availabilityPct')]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=c, value=v); cell.border = border
            if isinstance(v, float):
                cell.number_format = '0.00'
    for c, w in enumerate([18, 36, 8, 8, 10, 10, 10, 10, 10, 14], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    c = data.get('counts') or {}
    r = 8 + len(data['series'])
    ws.cell(row=r, column=1, value='Alarm events in the period').font = Font(bold=True)
    ws.cell(row=r + 1, column=1, value='Critical'); ws.cell(row=r + 1, column=2, value=c.get('critical', 0))
    ws.cell(row=r + 2, column=1, value='Warning');  ws.cell(row=r + 2, column=2, value=c.get('warning', 0))
    ws.cell(row=r + 3, column=1, value='Info');     ws.cell(row=r + 3, column=2, value=c.get('info', 0))

    # ── Data (wide: timestamp, then avg/min/max per tag) ──
    wd = wb.create_sheet('Data')
    cols = ['Timestamp']
    for s in data['series']:
        u = f" ({s['unit']})" if s.get('unit') else ''
        cols += [f"{s['tag']} avg{u}", f"{s['tag']} min{u}", f"{s['tag']} max{u}"]
    header(wd, 1, cols)
    times = sorted({p[0] for s in data['series'] for p in s['points']})
    lookup = [{p[0]: p for p in s['points']} for s in data['series']]
    for r_i, t in enumerate(times, start=2):
        wd.cell(row=r_i, column=1, value=parse_ts(t).replace(tzinfo=None)).number_format = 'yyyy-mm-dd hh:mm:ss'
        for s_i, m in enumerate(lookup):
            p = m.get(t)
            for k in range(3):
                v = p[1 + k] if p else None
                cell = wd.cell(row=r_i, column=2 + s_i * 3 + k, value=v)
                if v is not None:
                    cell.number_format = '0.000'
    wd.column_dimensions['A'].width = 20
    for c_i in range(2, len(cols) + 1):
        wd.column_dimensions[get_column_letter(c_i)].width = 16
    wd.freeze_panes = 'B2'

    # One line chart per tag on the summary sheet (average only).
    if times:
        anchor_row = r + 6
        for s_i, s in enumerate(data['series']):
            ch = LineChart()
            ch.title = f"{s['tag']} — {s.get('name') or ''}"
            ch.y_axis.title = s.get('unit') or ''
            ch.height = 6; ch.width = 18
            col = 2 + s_i * 3
            ch.add_data(Reference(wd, min_col=col, min_row=1, max_row=len(times) + 1), titles_from_data=True)
            ch.set_categories(Reference(wd, min_col=1, min_row=2, max_row=len(times) + 1))
            ch.legend = None
            ws.add_chart(ch, f'A{anchor_row + s_i * 13}')

    # ── Alarms ──
    wa = wb.create_sheet('Alarms')
    header(wa, 1, ['Raised', 'Cleared', 'Severity', 'State', 'Rule', 'Message', 'Acknowledged'])
    for i, e in enumerate(data.get('events') or [], start=2):
        vals = [parse_ts(e['triggeredAt']).replace(tzinfo=None) if e.get('triggeredAt') else None,
                parse_ts(e['clearedAt']).replace(tzinfo=None) if e.get('clearedAt') else None,
                e.get('severity'), e.get('state'), e.get('rule'), e.get('message'), 'yes' if e.get('acknowledged') else 'no']
        for c_i, v in enumerate(vals, start=1):
            cell = wa.cell(row=i, column=c_i, value=v)
            if c_i in (1, 2) and v is not None:
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
    for c_i, w in enumerate([20, 20, 10, 10, 40, 70, 12], start=1):
        wa.column_dimensions[get_column_letter(c_i)].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def main():
    data = json.load(sys.stdin)
    fmt = (data.get('format') or 'pdf').lower()
    blob = build_xlsx(data) if fmt == 'xlsx' else build_pdf(data)
    sys.stdout.buffer.write(blob)
    sys.stdout.buffer.flush()


if __name__ == '__main__':
    main()
